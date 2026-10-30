"""
Builder script to generate the complete B.Tech Project Report for
'AI Legal Contract Risk Analyzer' adhering strictly to IV/IV B.Tech guidelines.
"""
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_run_font(run, name="Times New Roman", size=12, bold=False, italic=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), name)
    rFonts.set(qn('w:hAnsi'), name)
    rFonts.set(qn('w:cs'), name)
    rPr.append(rFonts)

def add_p(doc, text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6, space_before=0, line_spacing=1.5, bold=False, italic=False, size=12):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = line_spacing
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold, italic=italic)
    return p

def add_bullet_p(doc, text="", space_after=4, space_before=0, line_spacing=1.5, size=12):
    p = doc.add_paragraph(style='List Bullet')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = line_spacing
    r = p.add_run(text)
    set_run_font(r, size=size)
    return p

def add_chap_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(18)
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(text.upper())
    set_run_font(r, size=14, bold=True)
    return p

def add_sec_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(text.upper())
    set_run_font(r, size=12, bold=True)
    return p

def add_subsec_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(text)
    set_run_font(r, size=12, bold=True)
    return p

def add_fig(doc, img_path, caption, width=Inches(5.5)):
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(6)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=width)
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(2)
        p_cap.paragraph_format.space_after = Pt(12)
        r_cap = p_cap.add_run(caption)
        set_run_font(r_cap, size=11, bold=True)
    else:
        print(f"Warning: image {img_path} not found!")

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(r'''
        <w:tblBorders %s>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="2C3E50"/>
            <w:left w:val="single" w:sz="6" w:space="0" w:color="2C3E50"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="2C3E50"/>
            <w:right w:val="single" w:sz="6" w:space="0" w:color="2C3E50"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
        </w:tblBorders>
    ''' % nsdecls('w'))
    tblPr.append(borders)

def add_custom_table(doc, headers, data, col_widths=None, caption=None):
    if caption:
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(8)
        p_cap.paragraph_format.space_after = Pt(4)
        r_cap = p_cap.add_run(caption)
        set_run_font(r_cap, size=11, bold=True)

    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    # Header row
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        shading = parse_xml(r'<w:shd %s w:fill="F1F5F9"/>' % nsdecls('w'))
        hdr_cells[i]._tc.get_or_add_tcPr().append(shading)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        for r in p.runs:
            set_run_font(r, size=10, bold=True)

    # Data rows
    for r_idx, row in enumerate(data):
        row_cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            row_cells[c_idx].text = str(val)
            row_cells[c_idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx > 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            for r in p.runs:
                set_run_font(r, size=9.5)

    if col_widths:
        for r in table.rows:
            for idx, w in enumerate(col_widths):
                r.cells[idx].width = w

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_after = Pt(8)
    return table

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    cell.width = Inches(5.77)
    
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(r'<w:shd %s w:fill="F8FAFC"/>' % nsdecls('w'))
    tcPr.append(shd)
    
    borders = parse_xml(r'''
        <w:tcBorders %s>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
            <w:left w:val="single" w:sz="18" w:space="0" w:color="2563EB"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>
        </w:tcBorders>
    ''' % nsdecls('w'))
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(code_text.strip())
    set_run_font(r, name="Consolas", size=9.5, bold=False)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_after = Pt(4)

def add_toc_entry(doc, title, page_no, level=0):
    p = doc.add_paragraph()
    p.paragraph_format.tab_stops.add_tab_stop(Inches(5.77), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_before = Pt(3 if level == 0 else 1)
    p.paragraph_format.space_after = Pt(1)
    
    if level == 0:
        p.paragraph_format.left_indent = Inches(0)
        r_t = p.add_run(title.upper())
        set_run_font(r_t, name="Times New Roman", size=11, bold=True)
        r_p = p.add_run(f"\t{page_no}")
        set_run_font(r_p, name="Times New Roman", size=11, bold=True)
    elif level == 1:
        p.paragraph_format.left_indent = Inches(0.2)
        r_t = p.add_run(title)
        set_run_font(r_t, name="Times New Roman", size=10.5, bold=False)
        r_p = p.add_run(f"\t{page_no}")
        set_run_font(r_p, name="Times New Roman", size=10.5, bold=False)
    else:
        p.paragraph_format.left_indent = Inches(0.4)
        r_t = p.add_run(title)
        set_run_font(r_t, name="Times New Roman", size=10, bold=False)
        r_p = p.add_run(f"\t{page_no}")
        set_run_font(r_p, name="Times New Roman", size=10, bold=False)
    return p

print("Base builder helpers loaded successfully.")
