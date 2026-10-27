import os
from io import BytesIO

from fastapi import (
    APIRouter,
    Depends,
    File,
    HTTPException,
    UploadFile,
)
from fastapi.responses import StreamingResponse
from groq import Groq
from pydantic import BaseModel, Field
from PyPDF2 import PdfReader
from docx import Document
import pymupdf as fitz
try:
    import easyocr
except ImportError:
    easyocr = None
import numpy as np
from PIL import Image
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import (
    getSampleStyleSheet,
    ParagraphStyle,
)
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
)
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from docx.shared import Pt

from services.auth_service import get_current_user


router = APIRouter(
    prefix="/api/language",
    tags=["Multi-language Support"],
    dependencies=[Depends(get_current_user)],
)


SUPPORTED_LANGUAGES = {
    # Indian languages in the Eighth Schedule
    "Assamese",
    "Bengali",
    "Bodo",
    "Dogri",
    "Gujarati",
    "Hindi",
    "Kannada",
    "Kashmiri",
    "Konkani",
    "Maithili",
    "Malayalam",
    "Manipuri",
    "Marathi",
    "Nepali",
    "Odia",
    "Punjabi",
    "Sanskrit",
    "Santali",
    "Sindhi",
    "Tamil",
    "Telugu",
    "Urdu",

    # Project interface language
    "English",

    # Additional international languages
    "Spanish",
    "French",
    "German",
    "Arabic",
}


class TranslationRequest(BaseModel):
    text: str = Field(
        min_length=1,
        max_length=20000
    )

    target_language: str = Field(
        min_length=2,
        max_length=30
    )


_ocr_reader = None

def get_ocr_reader():
    global _ocr_reader

    if easyocr is None:
        raise HTTPException(
            status_code=503,
            detail="OCR is unavailable because EasyOCR is not installed."
        )

    if _ocr_reader is None:
        _ocr_reader = easyocr.Reader(
            ["en"],
            gpu=False
        )

    return _ocr_reader


def extract_pdf_text_with_ocr(
    content: bytes
) -> str:
    try:
        pdf = fitz.open(
            stream=content,
            filetype="pdf"
        )

        reader = get_ocr_reader()
        pages = []

        for page in pdf:
            pix = page.get_pixmap(
                matrix=fitz.Matrix(2, 2),
                alpha=False
            )

            image = Image.frombytes(
                "RGB",
                [pix.width, pix.height],
                pix.samples
            )

            image_array = np.array(image)

            text_parts = reader.readtext(
                image_array,
                detail=0,
                paragraph=True
            )

            page_text = "\n".join(
                str(item).strip()
                for item in text_parts
                if str(item).strip()
            )

            if page_text:
                pages.append(page_text)

        pdf.close()

        return "\n\n".join(pages).strip()

    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail=(
                "The PDF could not be processed with OCR: "
                f"{exc}"
            ),
        ) from exc


def extract_uploaded_text(
    file_name: str,
    content: bytes
) -> str:

    extension = os.path.splitext(
        file_name.lower()
    )[1]

    if extension == ".pdf":
        try:
            reader = PdfReader(
                BytesIO(content)
            )

            pages = []

            for page in reader.pages:
                page_text = (
                    page.extract_text()
                    or ""
                )

                if page_text.strip():
                    pages.append(
                        page_text.strip()
                    )

            extracted = "\n\n".join(
                pages
            ).strip()

            # Automatic OCR fallback for
            # scanned/image-only PDFs.
            if not extracted:
                extracted = extract_pdf_text_with_ocr(
                    content
                )

            return extracted

        except Exception as exc:
            # If normal PDF parsing fails,
            # still try OCR.
            try:
                return extract_pdf_text_with_ocr(
                    content
                )

            except HTTPException:
                raise HTTPException(
                    status_code=400,
                    detail=(
                        f"Could not read the PDF file: "
                        f"{exc}"
                    ),
                ) from exc

    if extension == ".docx":
        try:
            document = Document(
                BytesIO(content)
            )

            paragraphs = [
                p.text.strip()
                for p in document.paragraphs
                if p.text.strip()
            ]

            return "\n\n".join(
                paragraphs
            ).strip()

        except Exception as exc:
            raise HTTPException(
                status_code=400,
                detail=(
                    f"Could not read the Word file: "
                    f"{exc}"
                ),
            ) from exc

    if extension == ".txt":
        try:
            return content.decode(
                "utf-8"
            ).strip()

        except UnicodeDecodeError as exc:
            raise HTTPException(
                status_code=400,
                detail=(
                    "The TXT file must use "
                    "UTF-8 text encoding."
                ),
            ) from exc

    raise HTTPException(
        status_code=400,
        detail=(
            "Unsupported file type. Please upload "
            "a PDF, Word (.docx), or TXT file."
        ),
    )


class TranslationDownloadRequest(BaseModel):
    translated_text: str = Field(
        min_length=1,
        max_length=50000
    )

    target_language: str = Field(
        min_length=2,
        max_length=30
    )


def safe_download_name(
    language: str
) -> str:

    cleaned = "".join(
        ch
        for ch in language
        if ch.isalnum()
        or ch in " _-"
    ).strip()

    return cleaned or "translated"


def add_translated_paragraphs(
    story,
    text: str,
    body_style
):
    for block in text.replace(
        "\r\n",
        "\n"
    ).split("\n"):

        line = block.strip()

        if not line:
            story.append(
                Spacer(1, 8)
            )
            continue

        escaped = (
            line
            .replace("&", "&amp;")
            .replace("<", "&lt;")
            .replace(">", "&gt;")
        )

        story.append(
            Paragraph(
                escaped,
                body_style
            )
        )

        story.append(
            Spacer(1, 5)
        )


FONT_BY_LANGUAGE = {
    "Assamese":
        "/usr/share/fonts/truetype/noto/NotoSansBengali-Regular.ttf",

    "Bengali":
        "/usr/share/fonts/truetype/noto/NotoSansBengali-Regular.ttf",

    "Bodo":
        "/usr/share/fonts/truetype/noto/NotoSansDevanagari-Regular.ttf",

    "Dogri":
        "/usr/share/fonts/truetype/noto/NotoSansDevanagari-Regular.ttf",

    "Gujarati":
        "/usr/share/fonts/truetype/noto/NotoSansGujarati-Regular.ttf",

    "Hindi":
        "/usr/share/fonts/truetype/noto/NotoSansDevanagari-Regular.ttf",

    "Kannada":
        "/usr/share/fonts/truetype/noto/NotoSansKannada-Regular.ttf",

    "Kashmiri":
        "/usr/share/fonts/truetype/noto/NotoSansArabic-Regular.ttf",

    "Konkani":
        "/usr/share/fonts/truetype/noto/NotoSansDevanagari-Regular.ttf",

    "Maithili":
        "/usr/share/fonts/truetype/noto/NotoSansDevanagari-Regular.ttf",

    "Malayalam":
        "/usr/share/fonts/truetype/noto/NotoSansMalayalam-Regular.ttf",

    "Manipuri":
        "/usr/share/fonts/truetype/noto/NotoSansBengali-Regular.ttf",

    "Marathi":
        "/usr/share/fonts/truetype/noto/NotoSansDevanagari-Regular.ttf",

    "Nepali":
        "/usr/share/fonts/truetype/noto/NotoSansDevanagari-Regular.ttf",

    "Odia":
        "/usr/share/fonts/truetype/noto/NotoSansOriya-Regular.ttf",

    "Punjabi":
        "/usr/share/fonts/truetype/noto/NotoSansGurmukhi-Regular.ttf",

    "Sanskrit":
        "/usr/share/fonts/truetype/noto/NotoSansDevanagari-Regular.ttf",

    "Santali":
        "/usr/share/fonts/truetype/noto/NotoSansOlChiki-Regular.ttf",

    "Sindhi":
        "/usr/share/fonts/truetype/noto/NotoSansArabic-Regular.ttf",

    "Tamil":
        "/usr/share/fonts/truetype/noto/NotoSansTamil-Regular.ttf",

    "Telugu":
        "/usr/share/fonts/truetype/noto/NotoSansTelugu-Regular.ttf",

    "Urdu":
        "/usr/share/fonts/truetype/noto/NotoSansArabic-Regular.ttf",

    "Arabic":
        "/usr/share/fonts/truetype/noto/NotoSansArabic-Regular.ttf",

    "English":
        "/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf",

    "Spanish":
        "/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf",

    "French":
        "/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf",

    "German":
        "/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf",
}


def get_pdf_font(
    language: str
) -> str:

    font_path = FONT_BY_LANGUAGE.get(
        language,
        "/usr/share/fonts/truetype/noto/NotoSans-Regular.ttf"
    )

    font_name = (
        f"LegalLens_"
        f"{language.replace('-', '_').replace(' ', '_')}"
    )

    try:
        pdfmetrics.getFont(
            font_name
        )

    except KeyError:
        try:
            pdfmetrics.registerFont(
                TTFont(
                    font_name,
                    font_path
                )
            )

        except Exception:
            return "Helvetica"

    return font_name


@router.post(
    "/translate-file"
)
async def translate_file(
    file: UploadFile = File(...)
):
    allowed_types = {
        ".pdf",
        ".docx",
        ".txt"
    }

    file_name = (
        file.filename
        or "uploaded_file"
    )

    extension = os.path.splitext(
        file_name.lower()
    )[1]

    if extension not in allowed_types:
        raise HTTPException(
            status_code=400,
            detail=(
                "Unsupported file type. Please upload "
                "a PDF, Word (.docx), or TXT file."
            )
        )

    content = await file.read()

    if not content:
        raise HTTPException(
            status_code=400,
            detail="The uploaded file is empty."
        )

    # Keep upload bounded to avoid
    # unexpectedly large translation requests.
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(
            status_code=400,
            detail=(
                "The file is too large. Please "
                "upload a file smaller than 10 MB."
            ),
        )

    extracted_text = extract_uploaded_text(
        file_name,
        content
    )

    if not extracted_text:
        raise HTTPException(
            status_code=400,
            detail=(
                "No readable text was found in this file. "
                "Please upload a clearer document or "
                "use the Scan feature."
            ),
        )

    # Keep translation endpoint and existing
    # text area limits unchanged.
    if len(extracted_text) > 20000:
        extracted_text = (
            extracted_text[:20000]
        )

    return {
        "file_name": file_name,
        "extracted_text": extracted_text,
        "message": (
            "Document text extracted successfully."
        ),
    }


@router.post(
    "/translate"
)
def translate_text(
    payload: TranslationRequest
):
    if payload.target_language not in SUPPORTED_LANGUAGES:
        raise HTTPException(
            status_code=400,
            detail=(
                "Unsupported language. Choose one of: "
                + ", ".join(
                    sorted(
                        SUPPORTED_LANGUAGES
                    )
                )
                + "."
            ),
        )

    api_key = os.getenv(
        "GROQ_API_KEY"
    )

    if not api_key:
        raise HTTPException(
            status_code=500,
            detail=(
                "Groq API configuration is missing."
            )
        )

    prompt = f"""
You are a professional multilingual legal-document translation assistant.

Translate the following text into {payload.target_language}.

Important instructions:
- Preserve the original meaning.
- Do not add new information.
- Do not remove important information.
- Preserve headings, numbered sections, bullet points, and paragraph structure where possible.
- Use clear and natural language.
- Keep legal terminology accurate and understandable.
- Return only the translated text.

TEXT TO TRANSLATE:
{payload.text}
"""

    try:
        client = Groq(
            api_key=api_key
        )

        response = client.chat.completions.create(
            model="openai/gpt-oss-120b",
            messages=[
                {
                    "role": "user",
                    "content": prompt,
                }
            ],
            temperature=0.1,
        )

        translated_text = (
            response.choices[0]
            .message.content
        )

        if not translated_text:
            raise HTTPException(
                status_code=500,
                detail=(
                    "No translated content was generated."
                )
            )

        return {
            "target_language":
                payload.target_language,
            "translated_text":
                translated_text,
        }

    except HTTPException:
        raise

    except Exception as exc:
        raise HTTPException(
            status_code=500,
            detail=(
                f"Translation failed: {exc}"
            ),
        ) from exc


@router.post(
    "/download-pdf"
)
def download_translated_pdf(
    payload: TranslationDownloadRequest
):
    if payload.target_language not in SUPPORTED_LANGUAGES:
        raise HTTPException(
            status_code=400,
            detail="Unsupported language."
        )

    buffer = BytesIO()

    styles = getSampleStyleSheet()

    font_name = get_pdf_font(
        payload.target_language
    )

    title_style = ParagraphStyle(
        "TranslatedTitle",
        parent=styles["Title"],
        alignment=TA_CENTER,
        fontSize=16,
        leading=20,
        fontName=font_name,
        spaceAfter=16,
    )

    body_style = ParagraphStyle(
        "TranslatedBody",
        parent=styles["BodyText"],
        fontName=font_name,
        fontSize=10.5,
        leading=15,
        spaceAfter=6,
    )

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=45,
        leftMargin=45,
        topMargin=45,
        bottomMargin=45,
        title=(
            "Translated Contract - "
            f"{payload.target_language}"
        ),
    )

    story = [
        Paragraph(
            (
                "Translated Contract "
                f"({payload.target_language})"
            ),
            title_style,
        ),
        Paragraph(
            "Generated by LegalLens AI",
            body_style
        ),
        Spacer(1, 10),
    ]

    add_translated_paragraphs(
        story,
        payload.translated_text,
        body_style
    )

    doc.build(story)

    buffer.seek(0)

    filename = (
        "translated_contract_"
        f"{safe_download_name(payload.target_language)}"
        ".pdf"
    )

    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={
            "Content-Disposition":
                f'attachment; filename="{filename}"'
        },
    )


@router.post(
    "/download-docx"
)
def download_translated_docx(
    payload: TranslationDownloadRequest
):
    if payload.target_language not in SUPPORTED_LANGUAGES:
        raise HTTPException(
            status_code=400,
            detail="Unsupported language."
        )

    document = Document()

    title = document.add_paragraph()

    title.alignment = 1

    run = title.add_run(
        (
            "Translated Contract "
            f"({payload.target_language})"
        )
    )

    run.bold = True
    run.font.size = Pt(16)

    subtitle = document.add_paragraph(
        "Generated by LegalLens AI"
    )

    subtitle.alignment = 1

    for block in payload.translated_text.replace(
        "\r\n",
        "\n"
    ).split("\n"):
        document.add_paragraph(
            block
        )

    buffer = BytesIO()

    document.save(
        buffer
    )

    buffer.seek(0)

    filename = (
        "translated_contract_"
        f"{safe_download_name(payload.target_language)}"
        ".docx"
    )

    return StreamingResponse(
        buffer,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
        headers={
            "Content-Disposition":
                f'attachment; filename="{filename}"'
        },
    )