
from pathlib import Path

from PyPDF2 import PdfReader
from docx import Document


SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


def extract_pdf_pages(file_path: str) -> list[str]:
    """Return PDF text as a list, preserving one entry per page."""

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    try:
        reader = PdfReader(str(path))
        pages = []

        for page in reader.pages:
            text = page.extract_text() or ""
            pages.append(text.strip())

        if not any(pages):
            raise ValueError(
                "The PDF does not contain readable text. "
                "It may require OCR."
            )

        return pages

    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(
            f"Unable to read PDF file: {exc}"
        ) from exc


def extract_text_from_pdf(file_path: str) -> str:
    """Extract PDF text while preserving the existing string interface."""

    pages = extract_pdf_pages(file_path)
    return "\n\n".join(text for text in pages if text).strip()


def extract_text_from_docx(file_path: str) -> str:
    """Extract paragraphs and tables from a DOCX file."""

    try:
        document = Document(file_path)
        paragraphs = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                paragraphs.append(text)

        for table in document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]

                if cells:
                    paragraphs.append(" | ".join(cells))

        result = "\n\n".join(paragraphs).strip()

        if not result:
            raise ValueError(
                "The DOCX file does not contain readable text."
            )

        return result

    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(
            f"Unable to read DOCX file: {exc}"
        ) from exc


def normalize_evidence_text(text: str) -> str:
    """Normalize whitespace for reliable quotation matching."""

    return " ".join((text or "").split()).casefold()


def verify_clause_evidence(
    original_snippet: str,
    contract_text: str,
    file_path: str | None = None,
    file_type: str | None = None,
) -> dict:
    """
    Verify a quoted clause against extracted contract text.

    Page numbers are returned only when a matching quote is found
    in an individual PDF page. DOCX and pasted-text contracts do
    not receive invented page numbers.
    """

    snippet = normalize_evidence_text(original_snippet)
    full_text = normalize_evidence_text(contract_text)

    if not snippet:
        return {
            "evidence_verified": False,
            "page_number": None,
            "evidence_status": "missing_quote",
        }

    if snippet not in full_text:
        return {
            "evidence_verified": False,
            "page_number": None,
            "evidence_status": "quote_not_found",
        }

    extension = (file_type or "").lower().lstrip(".")

    if extension == "pdf" and file_path:
        try:
            pages = extract_pdf_pages(file_path)

            for page_number, page_text in enumerate(pages, start=1):
                if snippet in normalize_evidence_text(page_text):
                    return {
                        "evidence_verified": True,
                        "page_number": page_number,
                        "evidence_status": "verified",
                    }

            # The quote exists in stored text, but its page
            # could not be independently confirmed.
            return {
                "evidence_verified": True,
                "page_number": None,
                "evidence_status": "page_not_confirmed",
            }

        except (OSError, ValueError):
            return {
                "evidence_verified": True,
                "page_number": None,
                "evidence_status": "page_lookup_failed",
            }

    return {
        "evidence_verified": True,
        "page_number": None,
        "evidence_status": "verified_page_unavailable",
    }


def extract_contract_text(file_path: str) -> str:
    """Extract contract text based on file extension."""

    path = Path(file_path)
    extension = path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            "Unsupported file type. "
            "Only PDF and DOCX files are supported."
        )

    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    if extension == ".pdf":
        return extract_text_from_pdf(str(path))

    if extension == ".docx":
        return extract_text_from_docx(str(path))

    raise ValueError("Unsupported document type.")
