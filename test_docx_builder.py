import docx
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()
s0 = doc.sections[0]
s0.page_width = Inches(8.27)
s0.page_height = Inches(11.69)
s0.left_margin = Inches(1.5)
s0.right_margin = Inches(1.0)
s0.top_margin = Inches(1.0)
s0.bottom_margin = Inches(1.0)

# Section 0 footer: lower roman
sectPr0 = s0._sectPr
pgNumType0 = OxmlElement('w:pgNumType')
pgNumType0.set(qn('w:fmt'), 'lowerRoman')
pgNumType0.set(qn('w:start'), '1')
sectPr0.append(pgNumType0)

p = s0.footer.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run()
r.font.name = 'Times New Roman'
fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
r._r.append(fldChar1)
r._r.append(instrText)
r._r.append(fldChar2)
r._r.append(fldChar3)

doc.add_paragraph("Test Preliminary Page")
s1 = doc.add_section()
s1.page_width = Inches(8.27)
s1.page_height = Inches(11.69)
s1.left_margin = Inches(1.5)
s1.right_margin = Inches(1.0)
s1.top_margin = Inches(1.0)
s1.bottom_margin = Inches(1.0)

sectPr1 = s1._sectPr
pgNumType1 = OxmlElement('w:pgNumType')
pgNumType1.set(qn('w:fmt'), 'decimal')
pgNumType1.set(qn('w:start'), '1')
sectPr1.append(pgNumType1)

p1 = s1.footer.paragraphs[0]
p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
r1 = p1.add_run()
r1.font.name = 'Times New Roman'
r1._r.append(parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w')))
r1._r.append(parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w')))
r1._r.append(parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w')))
r1._r.append(parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w')))

doc.add_paragraph("Test Main Page 1")
doc.save("test_page_num.docx")
print("Page number test succeeded!")
